"""Generación de memorias de cálculo Word para análisis de socavación.

Formato formal uniforme (Arial Narrow 11 pt, justificado, A4) con ecuaciones
nativas de Word (OMML): las fórmulas se renderizan como expresiones
matemáticas reales, nunca como funciones de texto.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt

from socavacion.domain.models import Proyecto
from socavacion.domain.results import ResultadoCompleto
from socavacion.report import docx_calcs, docx_sections
from socavacion.report.docx_format import (
    body,
    bottom_border,
    configure_document,
    configure_header_footer,
    enforce_uniform_typography,
    format_run,
    table,
)
from socavacion.report.docx_style import FILENAME_PREFIX, TEAL

REF_MTC = "Manual de Puentes MTC 2018"


def generate_socavacion_docx(
    resultado: ResultadoCompleto,
    proyecto: Proyecto,
    output_path: str | Path,
) -> Path:
    """Crea la memoria Word y devuelve su ruta absoluta."""
    path = Path(output_path).expanduser().resolve()
    if path.suffix.lower() != ".docx":
        path = path.with_suffix(".docx")
    path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    configure_document(document)
    configure_header_footer(document)

    _cover(document, resultado, proyecto)
    _contents(document)
    docx_sections.bases(document)
    docx_sections.datos_entrada(document, proyecto)
    docx_sections.caudales(document, resultado, proyecto)
    docx_sections.regimen(document, resultado, proyecto)
    docx_calcs.socavacion_general(document, resultado, proyecto)
    docx_calcs.socavacion_contraccion(document, resultado, proyecto)
    docx_calcs.socavacion_local(document, resultado, proyecto)
    docx_calcs.socavacion_total(document, resultado)
    docx_sections.geotecnia(document, resultado)
    docx_sections.referencias(document)

    enforce_uniform_typography(document)
    document.save(path)
    return path


def select_socavacion_docx_save_path(
    resultado: ResultadoCompleto | None = None,
    proyecto: Proyecto | None = None,
    *,
    initial_dir: Path | None = None,
) -> Path | None:
    """Abre diálogo nativo para guardar la memoria Word."""
    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError as exc:
        raise RuntimeError("Tkinter no está disponible para elegir el destino Word.") from exc

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    default_name = f"{FILENAME_PREFIX}_{datetime.now():%Y%m%d_%H%M}.docx"
    kwargs = {
        "parent": root,
        "title": "Guardar memoria de cálculo de socavación",
        "defaultextension": ".docx",
        "initialfile": default_name,
        "filetypes": (("Documento de Word", "*.docx"),),
    }
    if initial_dir is not None:
        kwargs["initialdir"] = str(initial_dir)
    try:
        selected = filedialog.asksaveasfilename(**kwargs)
    finally:
        root.destroy()
    return Path(selected) if selected else None


def generate_socavacion_docx_with_dialog(
    resultado: ResultadoCompleto,
    proyecto: Proyecto,
    *,
    initial_dir: Path | None = None,
) -> Path | None:
    """Solicita destino y genera la memoria Word."""
    path = select_socavacion_docx_save_path(resultado, proyecto, initial_dir=initial_dir)
    if path is None:
        print("Generación de la memoria Word cancelada por el usuario.")
        return None
    generated = generate_socavacion_docx(resultado, proyecto, path)
    print(f"Memoria Word guardada en: {generated}")
    return generated


def _cover(document: Document, resultado: ResultadoCompleto, proyecto: Proyecto) -> None:
    par = document.add_paragraph()
    par.paragraph_format.space_after = Pt(34)
    run = par.add_run("INGENIERÍA ESTRUCTURAL")
    format_run(run, 11, bold=True)
    bottom_border(par, TEAL, 16)

    par = document.add_paragraph(style="Title")
    par.add_run("Memoria de cálculo\nde socavación en puentes")

    par = document.add_paragraph(style="Subtitle")
    par.add_run(
        "Socavación general · contracción · local en estribos · cotas de cimentación · "
        "compatibilidad geotécnica"
    )

    par = document.add_paragraph()
    par.paragraph_format.space_before = Pt(32)
    par.paragraph_format.space_after = Pt(4)
    run = par.add_run("DOCUMENTO TÉCNICO DE DISEÑO")
    format_run(run, 11, bold=True)

    table(
        document,
        ("Dato", "Descripción"),
        (
            ("Proyecto", proyecto.nombre),
            ("Q diseño socavación", f"{resultado.Q_diseno:.2f} m³/s"),
            ("Q verificación socavación", f"{resultado.Q_verif:.2f} m³/s"),
            ("Norma principal", f"{REF_MTC} (HEC-18 / HEC-20)"),
            ("Alcance", "Socavación general, por contracción y local en estribos"),
            ("Fecha de emisión", datetime.now().strftime("%d/%m/%Y")),
        ),
        widths=(42, 125),
        accent=True,
    )

    par = document.add_paragraph()
    par.paragraph_format.space_before = Pt(30)
    run = par.add_run(
        "La presente memoria desarrolla las expresiones, leyendas, sustituciones numéricas, "
        "resultados y criterios técnicos del análisis de socavación según el Manual de "
        "Puentes MTC 2018 y los métodos HEC-18 / HEC-20."
    )
    format_run(run, 11, italic=True)
    document.add_page_break()


def _contents(document: Document) -> None:
    document.add_heading("Contenido", level=1)
    body(
        document,
        "La memoria se organiza siguiendo la secuencia del procedimiento de cálculo. Los "
        "títulos conservan estilos jerárquicos de Word para permitir navegación desde el "
        "panel de títulos.",
    )
    table(
        document,
        ("Sección", "Contenido desarrollado"),
        (
            ("1", "Bases de diseño y componentes evaluados"),
            ("2", "Datos de entrada del proyecto"),
            ("3", "Caudales de cálculo"),
            ("4", "Régimen del lecho"),
            ("5", "Socavación general"),
            ("6", "Socavación por contracción"),
            ("7", "Socavación local en estribos"),
            ("8", "Socavación total y cimentación"),
            ("9", "Compatibilidad geológico-geotécnica"),
            ("10", "Referencias normativas"),
        ),
        widths=(25, 142),
        accent=True,
    )
    document.add_page_break()
