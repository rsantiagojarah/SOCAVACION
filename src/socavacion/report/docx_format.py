"""Formato del documento Word: estilos, tablas, bloques de cálculo y ayudantes."""

from __future__ import annotations

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from socavacion.report.docx_math import (
    add_native_equation,
    es_expresion,
    legend_items,
    substitution_lines,
)
from socavacion.report.docx_style import (
    COVER_TITLE_FONT_SIZE_PT,
    FONT,
    GRAY,
    GREEN,
    HEADER_TEXT,
    LIGHT,
    MAX_CONTENT_WIDTH_MM,
    NAVY,
    NORMAL_MARGIN_MM,
    RED,
    RULE,
    TEAL,
    TEXT_COLOR,
    FOOTER_TEXT,
    UNIFORM_FONT_SIZE_PT,
)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(NORMAL_MARGIN_MM)
    section.bottom_margin = Mm(NORMAL_MARGIN_MM)
    section.left_margin = Mm(NORMAL_MARGIN_MM)
    section.right_margin = Mm(NORMAL_MARGIN_MM)
    section.header_distance = Mm(12.7)
    section.footer_distance = Mm(12.7)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    _font_style(normal, UNIFORM_FONT_SIZE_PT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, bold, before, after in (
        ("Title", 11, True, 0, 10),
        ("Subtitle", 11, False, 0, 8),
        ("Heading 1", 11, True, 14, 7),
        ("Heading 2", 11, True, 10, 4),
        ("Heading 3", 11, True, 7, 3),
    ):
        style = document.styles[name]
        _font_style(style, size, bold=bold)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.keep_with_next = True

    for name, italic in (("Caption", True), ("Intense Quote", False)):
        style = document.styles[name]
        _font_style(style, 11, italic=italic)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if "Formula" not in document.styles:
        style = document.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
        _font_style(style, 11)
        style.paragraph_format.left_indent = Mm(4)
        style.paragraph_format.right_indent = Mm(3)
        style.paragraph_format.space_before = Pt(1)
        style.paragraph_format.space_after = Pt(1)
    if "Equation" not in document.styles:
        style = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
        _font_style(style, 11)
        style.paragraph_format.left_indent = Mm(7)
        style.paragraph_format.right_indent = Mm(5)
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(4)
    bullet = document.styles["List Bullet"]
    _font_style(bullet, 11)
    bullet.paragraph_format.left_indent = Mm(11)
    bullet.paragraph_format.first_line_indent = Mm(-5)
    bullet.paragraph_format.space_after = Pt(2)
    bullet.paragraph_format.line_spacing = 1.5
    bullet.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _font_style(style, size: float, *, bold: bool = False, italic: bool = False) -> None:
    style.font.name = FONT
    style.font.size = Pt(UNIFORM_FONT_SIZE_PT)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def configure_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.clear()
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(HEADER_TEXT)
    format_run(run, 11, bold=True)
    bottom_border(header, RULE, 5)
    footer = section.footer.paragraphs[0]
    footer.clear()
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(FOOTER_TEXT)
    format_run(run, 11)
    field(footer, "PAGE")


def calc_block(
    document: Document,
    title: str,
    formula,
    legend: str,
    substitution: str,
    result: str,
    comment: str,
    reference: str,
) -> None:
    """Bloque formal: expresión OMML, leyenda, sustitución y conclusión."""
    heading = document.add_heading(title, level=3)
    heading.paragraph_format.keep_with_next = True
    formulas = (formula,) if isinstance(formula, str) else tuple(formula)
    intro = document.add_paragraph(style="Normal")
    intro.paragraph_format.keep_with_next = True
    if len(formulas) == 1:
        intro.add_run("Para desarrollar esta verificación se emplea la siguiente expresión:")
    else:
        intro.add_run("Para desarrollar esta verificación se emplean las siguientes expresiones:")
    for item in formulas:
        add_native_equation(document, item)

    where = document.add_paragraph(style="Normal")
    where.paragraph_format.space_before = Pt(1)
    where.paragraph_format.space_after = Pt(1)
    where.paragraph_format.keep_with_next = True
    run = where.add_run("Donde:")
    format_run(run, 11, bold=True)
    for definition in legend_items(legend):
        paragraph = document.add_paragraph(style="List Bullet")
        variable, separator, description = definition.partition(":")
        if separator:
            run = paragraph.add_run(variable.strip() + " = ")
            format_run(run, 11, bold=True)
            run = paragraph.add_run(description.strip())
            format_run(run, 11)
        else:
            run = paragraph.add_run(definition.strip())
            format_run(run, 11)

    replacing = document.add_paragraph(style="Normal")
    replacing.paragraph_format.space_before = Pt(3)
    replacing.paragraph_format.space_after = Pt(1)
    replacing.paragraph_format.keep_with_next = True
    run = replacing.add_run("Reemplazando los valores correspondientes:")
    format_run(run, 11)
    for line in substitution_lines(substitution):
        if es_expresion(line):
            add_native_equation(document, line)
            continue
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.left_indent = Mm(7)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(line)
        format_run(run, 11)

    conclusion = document.add_paragraph(style="Normal")
    conclusion.paragraph_format.space_before = Pt(2)
    conclusion.paragraph_format.space_after = Pt(3)
    run = conclusion.add_run("Por lo tanto, ")
    format_run(run, 11)
    run = conclusion.add_run(result)
    format_run(run, 11, bold=True)

    technical = document.add_paragraph(style="Normal")
    technical.paragraph_format.space_before = Pt(1)
    technical.paragraph_format.space_after = Pt(2)
    color = GREEN if "no cumple" not in comment.lower() else RED
    run = technical.add_run(comment)
    format_run(run, 11)

    citation = document.add_paragraph(style="Normal")
    citation.paragraph_format.space_before = Pt(0)
    citation.paragraph_format.space_after = Pt(7)
    run = citation.add_run("Referencia normativa: ")
    format_run(run, 11, bold=True)
    run = citation.add_run(reference)
    format_run(run, 11, italic=True)


def table(document, headers, rows, *, widths, accent=False) -> None:
    total_width = sum(float(width) for width in widths)
    if total_width > MAX_CONTENT_WIDTH_MM:
        scale = MAX_CONTENT_WIDTH_MM / total_width
        widths = tuple(float(width) * scale for width in widths)
    tbl = document.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER if accent else WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit = False
    tbl.style = "Table Grid"
    for cell, header, width in zip(tbl.rows[0].cells, headers, widths):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9D9D9")
        set_cell_margins(cell, 70, 80, 70, 80)
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.line_spacing = 1.5
        par.paragraph_format.keep_with_next = True
        run = par.add_run(str(header))
        format_run(run, 11, bold=True)
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if accent and row_index % 2:
                set_cell_shading(cell, LIGHT)
            set_cell_margins(cell, 55, 75, 55, 75)
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.line_spacing = 1.5
            run = par.add_run(str(value))
            format_run(run, 11)
    tbl.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for table_row in tbl.rows:
        row_properties = table_row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
    _set_table_geometry(tbl, widths, indent_twips=0 if not accent else 120)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _set_table_geometry(tbl, widths, *, indent_twips: int) -> None:
    width_twips = tuple(round(float(width) * 56.692913) for width in widths)
    total = sum(width_twips)
    properties = tbl._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total))
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(indent_twips))
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in width_twips:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in tbl.rows:
        for cell, width in zip(row.cells, width_twips):
            tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(width))


def body(document, text_value: str) -> None:
    par = document.add_paragraph(text_value)
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def comment(document, text_value: str) -> None:
    par = document.add_paragraph(style="Formula")
    shade_paragraph(par, LIGHT)
    left_border(par, TEAL, 12)
    run = par.add_run(text_value)
    format_run(run, 11)


def field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Actualizar campo"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend((begin, instr, separate, text, end))
    format_run(run, 11)


def format_run(run, size: float, *, bold: bool = False, italic: bool = False) -> None:
    run.font.name = FONT
    run.font.size = Pt(UNIFORM_FONT_SIZE_PT)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    fonts.set(qn("w:eastAsia"), FONT)


def enforce_uniform_typography(document: Document) -> None:
    paragraphs = list(document.paragraphs)
    for tbl in document.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    seen: set[int] = set()
    for paragraph in paragraphs:
        identity = id(paragraph._p)
        if identity in seen:
            continue
        seen.add(identity)
        is_cover_title = paragraph.style is not None and paragraph.style.name == "Title"
        paragraph.paragraph_format.line_spacing = 1.0 if is_cover_title else 1.5
        if paragraph.text.strip():
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if is_cover_title else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(
                COVER_TITLE_FONT_SIZE_PT if is_cover_title else UNIFORM_FONT_SIZE_PT
            )
            run.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
            fonts = run._element.get_or_add_rPr().rFonts
            fonts.set(qn("w:ascii"), FONT)
            fonts.set(qn("w:hAnsi"), FONT)
            fonts.set(qn("w:eastAsia"), FONT)


def bottom_border(paragraph, color: str, size: int) -> None:
    _paragraph_border(paragraph, "bottom", color, size)


def left_border(paragraph, color: str, size: int) -> None:
    _paragraph_border(paragraph, "left", color, size)


def _paragraph_border(paragraph, edge: str, color: str, size: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    existing = borders.find(qn(f"w:{edge}"))
    if existing is not None:
        borders.remove(existing)
    border = OxmlElement(f"w:{edge}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), color)
    borders.append(border)


def shade_paragraph(paragraph, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int, left: int, bottom: int, right: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
