"""Secciones de la memoria: bases, datos de entrada, caudales, régimen, geotecnia y referencias."""

from __future__ import annotations

from docx import Document

from socavacion.domain.enums import CondicionCaudal, LadoEstribo
from socavacion.domain.models import Estribo, Proyecto
from socavacion.domain.results import ResultadoCompleto, ResultadoCaudal
from socavacion.report.docx_format import body, calc_block, comment, table

LADO_LARGO = {LadoEstribo.IZQUIERDO: "izquierdo", LadoEstribo.DERECHO: "derecho"}
LADO_CORTO = {LadoEstribo.IZQUIERDO: "E-I", LadoEstribo.DERECHO: "E-D"}


def hidraulica_de(estribo: Estribo, cond: CondicionCaudal):
    return estribo.q100 if cond == CondicionCaudal.DISENO else estribo.q500


def etiqueta_caudal(caudal: ResultadoCaudal) -> str:
    if caudal.condicion == CondicionCaudal.DISENO:
        return f"Q100 (diseño) — Q = {caudal.Q:.1f} m³/s"
    return f"Q500 (verificación) — Q = {caudal.Q:.1f} m³/s"


def bases(document: Document) -> None:
    document.add_heading("1. Bases de diseño", level=1)
    body(
        document,
        "El cálculo se desarrolla conforme al Manual de Puentes MTC 2018 (Artículos 1.2 y "
        "2.4.3.8.3.4), complementado con los procedimientos HEC-18 y HEC-20 de la FHWA. Se "
        "evalúan los componentes de socavación general, por contracción y local en estribos "
        "para las condiciones de diseño Q100 y de verificación Q500, y se determinan las cotas "
        "mínimas de cimentación con la reserva reglamentaria.",
    )
    document.add_heading("1.1 Componentes evaluados", level=2)
    table(
        document,
        ("Componente", "Descripción", "Método"),
        (
            ("y_sg", "Socavación general (degradación del cauce)", "Lischtvan-Lebediev, Neill, Lacey"),
            ("y_sc", "Socavación por contracción", "Laursen (lecho vivo / agua clara)"),
            ("y_sl", "Socavación local en estribos", "Froehlich / HIRE, límites HEC-18"),
            ("y_s", "Socavación total", "Suma de componentes + largo plazo"),
        ),
        widths=(22, 95, 50),
    )
    comment(
        document,
        "Las tres componentes se calculan de forma independiente para cada estribo y cada "
        "condición de caudal; el diseño se rige por la envolvente de Q100 y Q500.",
    )


def datos_entrada(document: Document, proyecto: Proyecto) -> None:
    document.add_heading("2. Datos de entrada", level=1)

    document.add_heading("2.1 Caudales del proyecto", level=2)
    table(
        document,
        ("Parámetro", "Valor"),
        (
            ("Q100 (diseño)", f"{proyecto.Q100:.2f} m³/s"),
            ("Q500 (verificación)", f"{proyecto.Q500:.2f} m³/s"),
            (
                "Q desbordamiento",
                f"{proyecto.Q_ot:.2f} m³/s" if proyecto.Q_ot else "No ingresado",
            ),
        ),
        widths=(78, 89),
        accent=True,
    )

    document.add_heading("2.2 Clasificación del cauce", level=2)
    table(
        document,
        ("Parámetro", "Valor"),
        (
            ("Tipo de cauce", proyecto.cauce.tipo.value),
            ("Degradación de largo plazo y_sg,lp", f"{proyecto.cauce.y_sg_lp:.3f} m"),
        ),
        widths=(78, 89),
        accent=True,
    )

    for idx, estribo in enumerate(proyecto.estribos()):
        lado = LADO_LARGO[estribo.lado]
        document.add_heading(f"2.{idx + 3} Estribo {lado}", level=2)
        for etiqueta, hid in (("Q100", estribo.q100), ("Q500", estribo.q500)):
            table(
                document,
                (f"Hidráulica {etiqueta}", "Valor"),
                (
                    ("Tirante aproximación y1", f"{hid.y1:.3f} m"),
                    ("Velocidad aproximación V1", f"{hid.V1:.3f} m/s"),
                    ("Ancho cauce W1", f"{hid.W1:.2f} m"),
                    ("Luz hidráulica W2", f"{hid.W2:.2f} m"),
                    ("Tirante bajo puente y0", f"{hid.y0:.3f} m"),
                    ("Pendiente de energía Sf", f"{hid.Sf:.5f} m/m"),
                ),
                widths=(78, 89),
                accent=True,
            )
        table(
            document,
            ("Parámetro", "Valor"),
            (
                ("D50 granulometría", f"{estribo.D50_mm:.2f} mm"),
                ("Cota lecho Z", f"{estribo.Z_lecho:.3f} m s.n.m."),
                ("Forma del estribo", estribo.forma.value),
                ("Ángulo de ataque", f"{estribo.angulo_ataque:.1f}°"),
                ("Longitud L′", f"{estribo.L_prima:.2f} m"),
                ("Área obstruida Ae", f"{estribo.Ae:.2f} m²"),
            ),
            widths=(78, 89),
            accent=True,
        )

    document.add_heading("2.5 Geotecnia", level=2)
    table(
        document,
        ("Parámetro", "Valor"),
        (
            ("Cota fondo sondaje", f"{proyecto.geotecnia.cota_sondaje_min:.3f} m s.n.m."),
            ("Estrato competente", "Sí" if proyecto.geotecnia.hay_estrato_competente else "No"),
            ("Roca resistente", "Sí" if proyecto.geotecnia.roca_resistente else "No"),
            ("Tipo de cimentación", proyecto.geotecnia.tipo_cimentacion.value),
        ),
        widths=(78, 89),
        accent=True,
    )


def caudales(document: Document, resultado: ResultadoCompleto, proyecto: Proyecto) -> None:
    document.add_heading("3. Caudales de cálculo", level=1)
    body(
        document,
        "El caudal de diseño de socavación es el mayor entre Q100 y el caudal de "
        "desbordamiento; el de verificación es el mayor entre Q500 y el de desbordamiento "
        "(MTC 2018, Art. 1.2.3).",
    )
    if proyecto.Q_ot is not None:
        calc_block(
            document,
            "Selección de caudales de socavación",
            ("Q_{diseño} = max(Q_{100}, Q_{ot})", "Q_{verif} = max(Q_{500}, Q_{ot})"),
            "Q_{diseño}: caudal de diseño de socavación; Q_{verif}: caudal de verificación; "
            "Q_{100}, Q_{500}: caudales de retorno 100 y 500 años; Q_{ot}: caudal de desbordamiento",
            (
                f"Q_{{diseño}} = max({proyecto.Q100:.1f}, {proyecto.Q_ot:.1f}) = {resultado.Q_diseno:.1f} m³/s\n"
                f"Q_{{verif}} = max({proyecto.Q500:.1f}, {proyecto.Q_ot:.1f}) = {resultado.Q_verif:.1f} m³/s"
            ),
            f"se calcula con Q = {resultado.Q_diseno:.1f} m³/s y se verifica con Q = {resultado.Q_verif:.1f} m³/s.",
            "La envolvente con el caudal de desbordamiento cubre el escenario más desfavorable.",
            "Manual de Puentes MTC 2018, Art. 1.2.3.",
        )
    else:
        table(
            document,
            ("Condición", "Caudal (m³/s)"),
            (
                ("Diseño de socavación (Q100)", f"{resultado.Q_diseno:.2f}"),
                ("Verificación de socavación (Q500)", f"{resultado.Q_verif:.2f}"),
            ),
            widths=(95, 72),
            accent=True,
        )


def regimen(document: Document, resultado: ResultadoCompleto, proyecto: Proyecto) -> None:
    document.add_heading("4. Régimen del lecho", level=1)
    body(
        document,
        "El régimen del lecho se determina comparando la velocidad media de aproximación V1 "
        "con la velocidad crítica Vc de inicio de movimiento del sedimento. Si V1 supera a Vc "
        "existe transporte desde aguas arriba (lecho vivo); en caso contrario el flujo llega "
        "sin carga sólida (agua clara).",
    )
    estribos = {e.lado: e for e in proyecto.estribos()}
    for caudal in resultado.caudales:
        document.add_heading(f"4.{1 if caudal.condicion == CondicionCaudal.DISENO else 2} {etiqueta_caudal(caudal)}", level=2)
        for r in caudal.estribos:
            est = estribos[r.lado]
            hid = hidraulica_de(est, caudal.condicion)
            reg = r.regimen
            lado = LADO_LARGO[r.lado]
            es_vivo = reg.regimen.value == "lecho_vivo"
            ratio = reg.V_star / reg.omega if reg.omega > 0 else 0.0
            substitution = (
                f"V_c = 6.19 · {hid.y1:.3f}^(1/6) · {est.D50_m:.4f}^(1/3) = {reg.Vc:.3f} m/s\n"
                f"F_r = {reg.V1:.3f} / √(9.81 · {hid.y1:.3f}) = {reg.Fr:.3f}\n"
                f"V_* = √(9.81 · {hid.y1:.3f} · {hid.Sf:.5f}) = {reg.V_star:.4f} m/s\n"
                f"V_*/ω = {ratio:.3f}\n"
                f"k_1 = {reg.k1_contraccion:.3f}"
            )
            if es_vivo:
                result = (
                    f"V1 = {reg.V1:.3f} m/s > Vc = {reg.Vc:.3f} m/s; "
                    "el lecho trabaja en régimen de lecho vivo."
                )
                comment_text = (
                    f"Existe aporte de sedimento desde aguas arriba. El exponente de Laursen "
                    f"k1 = {reg.k1_contraccion:.3f} se obtiene de V*/ω = {ratio:.2f}."
                )
            else:
                result = (
                    f"V1 = {reg.V1:.3f} m/s ≤ Vc = {reg.Vc:.3f} m/s; "
                    "el lecho trabaja en régimen de agua clara."
                )
                comment_text = (
                    f"No hay aporte de sedimento desde aguas arriba. El exponente de Laursen "
                    f"k1 = {reg.k1_contraccion:.3f} se obtiene de V*/ω = {ratio:.2f}."
                )
            calc_block(
                document,
                f"Régimen del lecho — estribo {lado}",
                (
                    "V_c = K_u · y_1^(1/6) · D_{50}^(1/3)",
                    "F_r = V_1 / √(g · y_1)",
                    "V_* = √(g · y_1 · S_f)",
                ),
                "V_c: velocidad crítica de inicio de movimiento del sedimento; "
                "K_u: coeficiente de unidades, 6.19 en el sistema SI; "
                "y_1: tirante medio de aproximación; "
                "D_50: diámetro medio del material del lecho; "
                "F_r: número de Froude de aproximación; "
                "g: aceleración de la gravedad; "
                "V_*: velocidad de corte; "
                "S_f: pendiente de la línea de energía",
                substitution,
                result,
                comment_text,
                "FHWA HEC-18, Cap. 6 (velocidad crítica y régimen del lecho).",
            )


def geotecnia(document: Document, resultado: ResultadoCompleto) -> None:
    document.add_heading("9. Compatibilidad geológico-geotécnica", level=1)
    body(
        document,
        "Se verifica la compatibilidad entre la profundidad de socavación calculada y la "
        "información geotécnica disponible, según la matriz del documento interno 10.11.",
    )
    rows = []
    for item in resultado.geotecnia.items:
        ok = "Sí" if item.compatible else "No"
        rows.append((item.item, item.hidraulica, item.geotecnia, ok))
    table(
        document,
        ("Ítem", "Hidráulica", "Geotecnia", "¿OK?"),
        rows,
        widths=(42, 60, 60, 25),
        accent=True,
    )
    color_comment = resultado.geotecnia.conclusion
    comment(document, f"Conclusión: {color_comment}")


def tabla_resumen_totales(document: Document, resultado: ResultadoCompleto) -> None:
    rows = []
    for e in resultado.estribos_finales:
        lado = LADO_LARGO[e.lado]
        z_cim = f"{e.Z_cim_min:.3f}" if e.Z_cim_min is not None else "—"
        rows.append(
            (
                lado,
                f"{e.y_s_100:.3f}",
                f"{e.y_s_500:.3f}",
                f"{e.y_s_diseno:.3f}",
                f"{e.Z_lecho_actual:.3f}",
                f"{e.Z_lecho_soc:.3f}",
                z_cim,
            )
        )
    table(
        document,
        (
            "Estribo",
            "y_s Q100 (m)",
            "y_s Q500 (m)",
            "y_s diseño (m)",
            "Z lecho (m)",
            "Z soc (m)",
            "Z cim mín (m)",
        ),
        rows,
        widths=(24, 27, 27, 27, 27, 26, 26),
        accent=True,
    )


def referencias(document: Document) -> None:
    document.add_heading("10. Referencias normativas", level=1)
    for text in (
        "Manual de Puentes MTC 2018, Artículos 1.2, 1.2.3a, 1.2.4, 1.2.6 y 2.4.3.8.3.4.",
        "FHWA HEC-18, Evaluating Scour at Bridges: régimen del lecho, socavación por contracción (Laursen) y socavación local (Froehlich, HIRE).",
        "FHWA HEC-20, Stream Stability at Highway Structures.",
        "Lischtvan–Lebediev, Neill y Lacey: métodos clásicos de socavación general.",
        "Documento interno 10: Análisis de Socavación — Proceso de Cálculo (matriz 10.11).",
    ):
        document.add_paragraph(text, style="List Bullet")
